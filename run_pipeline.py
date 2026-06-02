# CLI runner
"""
run_pipeline.py
---------------
CLI entry point to run the migration pipeline directly.

Usage:
  python run_pipeline.py --input services.xlsx
  python run_pipeline.py --input services.xlsx --output results.xlsx
  python run_pipeline.py --create-sample  # Generate a sample input XLSX
"""

import argparse
import logging
import sys
from pathlib import Path

from dotenv import load_dotenv
load_dotenv()

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)

logger = logging.getLogger(__name__)


# ─────────────────────────────────────────────────────────────────────────────
# SAMPLE XLSX GENERATOR
# ─────────────────────────────────────────────────────────────────────────────

def create_sample_xlsx(output_path: str = "sample_services.xlsx"):
    """Create a sample input XLSX with example GCP and Azure services."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Services"

    headers = [
        "service_name", "service_type", "current_provider",
        "instance_type", "vcpus", "memory_gib",
        "region", "tenancy", "operating_system",
        "storage_gb", "number_of_instances",
        "current_monthly_cost_usd",
        "database_engine", "multi_az",
        "memory_mb", "invocations_per_month", "avg_duration_ms",
        "data_transfer_out_gb", "nat_gateways",
        "workload", "notes",
    ]

    # Header row styling
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri")
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = max(15, len(h) + 2)

    # Sample rows
    rows = [
        # GCP services
        ["Web Server",        "ec2", "GCP",   "n2-standard-4",   4,  16,   "us-central1",    "Shared",    "Linux",   100,  3,  450.00, None, None, None, None, None, None, None, "Constant usage", "Main web tier"],
        ["App Server",        "ec2", "GCP",   "n2-standard-8",   8,  32,   "us-central1",    "Shared",    "Linux",   50,   2,  380.00, None, None, None, None, None, None, None, "Variable usage", "Application layer"],
        ["DB Primary",        "rds", "GCP",   "n1-highmem-4",    4,  26,   "us-east1",       "Shared",    "Linux",   500,  1,  320.00, "MySQL", True, None, None, None, None, None, "Constant usage", "Production DB"],
        ["Object Store",      "s3",  "GCP",   None,              None, None, "us-central1",   "Shared",    None,      5000, 1,  100.00, None, None, None, None, None, 500, None, None, "GCS bucket ~5TB"],
        ["Serverless Fn",     "lambda","GCP", None,              None, None, "us-central1",   "Shared",    None,      None, 1,  30.00,  None, None, 512, 5000000, 200, None, None, None, "Cloud Functions"],
        # Azure services
        ["Cache Server",      "ec2", "Azure", "Standard_D4s_v3", 4,  16,   "eastus",         "Shared",    "Linux",   200,  2,  290.00, None, None, None, None, None, None, None, "Constant usage", "Redis cache hosts"],
        ["Analytics DB",      "rds", "Azure", "Standard_E8s_v3", 8,  64,   "eastus2",        "Shared",    "Linux",   1000, 1,  650.00, "PostgreSQL", False, None, None, None, None, None, "Variable usage", "OLAP workload"],
        ["API Gateway VNet",  "vpc", "Azure", None,              None, None, "westeurope",    "Shared",    None,      None, 1,  45.00,  None, None, None, None, None, 2000, 2, None, "Azure VNet + NAT"],
        ["ML Training VM",    "ec2", "Azure", "Standard_F16s_v2",16, 32,   "eastus",         "Dedicated", "Linux",   500,  1,  800.00, None, None, None, None, None, None, None, "Constant usage", "GPU-adjacent compute"],
        ["Logs Storage",      "s3",  "Azure", None,              None, None, "eastus",        "Shared",    None,      10000,1,  190.00, None, None, None, None, None, 100, None, None, "Azure Blob ~10TB"],
    ]

    data_fill_even = PatternFill("solid", fgColor="F2F2F2")
    data_font = Font(name="Calibri", size=10)

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = data_font
            if r_idx % 2 == 0:
                cell.fill = data_fill_even

    ws.freeze_panes = "A2"
    wb.save(output_path)
    print(f"✅ Sample XLSX created: {output_path}")
    print(f"   {len(rows)} services from GCP and Azure ready for estimation.")


# ─────────────────────────────────────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Cloud Migration Cost Estimator — map GCP/Azure to AWS"
    )
    parser.add_argument("--input",  "-i", help="Input XLSX file path")
    parser.add_argument("--output", "-o", help="Output XLSX file path (optional)")
    parser.add_argument("--create-sample", action="store_true",
                        help="Generate a sample input XLSX and exit")
    parser.add_argument("--sample-output", default="sample_services.xlsx",
                        help="Sample XLSX output path")

    args = parser.parse_args()

    if args.create_sample:
        create_sample_xlsx(args.sample_output)
        sys.exit(0)

    if not args.input:
        parser.print_help()
        sys.exit(1)

    if not Path(args.input).exists():
        print(f"❌ Input file not found: {args.input}")
        sys.exit(1)

    from agents.orchestrator import run_migration_pipeline

    print(f"\n🚀 Starting Cloud Migration Cost Estimation")
    print(f"   Input:  {args.input}")

    output = run_migration_pipeline(args.input, args.output)

    print(f"\n✅ Done! Output saved to: {output}")


if __name__ == "__main__":
    main()